"""
Text extraction, cleaning, and chunking for:

    .txt
    .pdf
    .docx

This module is used by both:

1. Permanent document processing
2. Temporary chat attachments

IMPORTANT:
This module itself does NOT:
- create database records
- save documents permanently
- create FAISS indexes
- modify the permanent document library

It only:
    file -> extracted text -> cleaned text -> chunks
"""

import logging
import re
from pathlib import Path
from typing import List, Union

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from app.core.config import settings
from app.services.ocr import ocr_page_image


logger = logging.getLogger(__name__)


# ============================================================
# OPTIONAL SENTENCE SPLITTER
# ============================================================

try:
    from sentence_splitter import SentenceSplitter

    _splitter = SentenceSplitter(
        language="en"
    )

except Exception:
    _splitter = None


# ============================================================
# TYPE
# ============================================================

FilePath = Union[str, Path]


# ============================================================
# TEXT EXTRACTION
# ============================================================

def extract_text_from_file(
    filepath: FilePath,
) -> str:
    """
    Extract readable text from TXT, PDF or DOCX.

    Supports both:
        str
        pathlib.Path

    Example:
        extract_text_from_file("resume.pdf")

    or:

        extract_text_from_file(Path("resume.pdf"))

    Returns:
        Extracted text as a string.

    Returns empty string if extraction fails.
    """

    # --------------------------------------------------------
    # Normalize Path / String
    # --------------------------------------------------------

    path = Path(filepath)

    if not path.exists():
        logger.warning(
            "File does not exist: %s",
            path,
        )
        return ""

    if not path.is_file():
        logger.warning(
            "Path is not a file: %s",
            path,
        )
        return ""

    extension = (
        path.suffix.lower()
    )

    try:

        # ====================================================
        # TXT
        # ====================================================

        if extension == ".txt":

            return _extract_txt(
                path
            )


        # ====================================================
        # PDF
        # ====================================================

        if extension == ".pdf":

            return _extract_pdf(
                path
            )


        # ====================================================
        # DOCX
        # ====================================================

        if extension == ".docx":

            return _extract_docx(
                path
            )


        # ====================================================
        # UNSUPPORTED
        # ====================================================

        logger.warning(
            "Unsupported file type for extraction: %s",
            path,
        )

        return ""

    except Exception:
        logger.exception(
            "Failed to extract text from %s",
            path,
        )

        return ""


# ============================================================
# TXT EXTRACTION
# ============================================================

def _extract_txt(
    path: Path,
) -> str:
    """
    Extract text from a TXT file.

    UTF-8 is attempted first.
    If the file contains invalid UTF-8 bytes,
    invalid characters are ignored.
    """

    try:

        with path.open(
            "r",
            encoding="utf-8",
            errors="ignore",
        ) as file:

            return file.read()

    except Exception:
        logger.exception(
            "Failed to read TXT file: %s",
            path,
        )

        return ""


# ============================================================
# PDF EXTRACTION
# ============================================================

def _extract_pdf(
    path: Path,
) -> str:
    """
    Extract text from a PDF using PyMuPDF.

    Each page is separated by a newline.

    This works well for normal text-based PDFs such as:
        - resumes
        - reports
        - research papers
        - assignments
        - documentation

    For pages where PyMuPDF finds NO text (i.e. the page is a
    scanned image), the page is rendered to a PNG and sent to
    a free OCR API as a fallback. This keeps normal text-based
    PDFs fast (no OCR calls) while still handling scanned pages.
    """

    page_texts = []

    ocr_pages_used = 0

    with fitz.open(
        str(path)
    ) as document:

        logger.info(
            "Extracting PDF: %s (%d pages)",
            path.name,
            len(document),
        )

        for page_number, page in enumerate(
            document,
            start=1,
        ):

            try:

                text = page.get_text(
                    "text"
                ) or ""

                text = text.strip()

                # ------------------------------------------------
                # OCR FALLBACK
                #
                # No extractable text usually means this page is
                # a scanned image rather than real text.
                # ------------------------------------------------

                if (
                    not text
                    and ocr_pages_used
                    < settings.OCR_MAX_PAGES_PER_DOCUMENT
                ):

                    logger.info(
                        "Page %d has no extractable text. "
                        "Attempting OCR.",
                        page_number,
                    )

                    pixmap = page.get_pixmap(
                        dpi=200
                    )

                    image_bytes = pixmap.tobytes(
                        "png"
                    )

                    ocr_text = ocr_page_image(
                        image_bytes
                    )

                    ocr_pages_used += 1

                    if ocr_text:

                        text = ocr_text.strip()

                        logger.info(
                            "OCR succeeded for page %d: "
                            "%d characters",
                            page_number,
                            len(text),
                        )

                    else:

                        logger.warning(
                            "OCR returned no text for "
                            "page %d.",
                            page_number,
                        )

                if text:
                    page_texts.append(
                        f"[Page {page_number}]\n{text}"
                    )

            except Exception:
                logger.exception(
                    "Failed to extract page %d from %s",
                    page_number,
                    path,
                )

    extracted = "\n\n".join(
        page_texts
    )

    if not extracted.strip():

        logger.warning(
            "No text could be extracted from PDF: %s "
            "even after OCR.",
            path,
        )

    return extracted


# ============================================================
# DOCX EXTRACTION
# ============================================================

def _extract_docx(
    path: Path,
) -> str:
    """
    Extract text from DOCX.

    Extracts:
        - paragraphs
        - table cells

    This is more reliable than extracting paragraphs alone
    because resumes and reports often contain information
    inside tables.
    """

    document = DocxDocument(
        str(path)
    )

    parts = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)


    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        table_parts = []

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                value = (
                    cell.text
                    .strip()
                )

                if value:
                    row_values.append(
                        value
                    )

            if row_values:
                table_parts.append(
                    " | ".join(
                        row_values
                    )
                )

        if table_parts:

            parts.append(
                f"[Table {table_index}]"
            )

            parts.extend(
                table_parts
            )


    return "\n".join(
        parts
    )


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(
    text: str,
) -> str:
    """
    Normalize extracted document text.

    Important:
    We intentionally preserve Unicode characters.

    This means names/content such as:
        é
        ü
        ñ
        भारतीय
        résumé

    are not unnecessarily deleted.
    """

    if not text:
        return ""

    # --------------------------------------------------------
    # Normalize non-breaking spaces
    # --------------------------------------------------------

    text = text.replace(
        "\u00A0",
        " ",
    )

    # --------------------------------------------------------
    # Normalize tabs/spaces
    #
    # Keep newlines because they preserve document structure.
    # --------------------------------------------------------

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    # --------------------------------------------------------
    # Remove excessive blank lines
    #
    # Example:
    #
    # A
    #
    #
    #
    # B
    #
    # becomes:
    #
    # A
    #
    # B
    # --------------------------------------------------------

    text = re.sub(
        r"\n\s*\n+",
        "\n\n",
        text,
    )

    # --------------------------------------------------------
    # Remove control characters except:
    #
    # \n
    # \r
    # \t
    # --------------------------------------------------------

    text = "".join(
        character
        for character in text
        if (
            character in "\n\r\t"
            or ord(character) >= 32
        )
    )

    # --------------------------------------------------------
    # Normalize Windows line endings
    # --------------------------------------------------------

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    return text.strip()


# ============================================================
# CHUNKING
# ============================================================

def chunk_text(
    text: str,
    chunk_size: int = 300,
    overlap: int = 50,
) -> List[str]:
    """
    Split document text into chunks.

    chunk_size:
        Approximate number of whitespace-separated tokens.

    overlap:
        Number of tokens repeated between consecutive chunks.

    Example:

        chunk 1:
        tokens 1-300

        chunk 2:
        tokens 251-550

        overlap = 50
    """

    if not text or not text.strip():
        return []

    # --------------------------------------------------------
    # Safety validation
    # --------------------------------------------------------

    if chunk_size <= 0:
        raise ValueError(
            "chunk_size must be greater than 0."
        )

    if overlap < 0:
        overlap = 0

    if overlap >= chunk_size:
        overlap = max(
            0,
            chunk_size // 5,
        )

    # --------------------------------------------------------
    # Clean before chunking
    # --------------------------------------------------------

    text = clean_text(
        text
    )

    if not text:
        return []

    # ========================================================
    # SENTENCE SPLITTING
    # ========================================================

    sentences = None

    if _splitter:

        try:

            sentences = _splitter.split(
                text
            )

        except Exception:
            logger.warning(
                "Sentence splitter failed. "
                "Falling back to line/sentence splitting."
            )

            sentences = None


    # ========================================================
    # FALLBACK SPLITTER
    # ========================================================

    if not sentences or len(sentences) <= 1:

        parts = []

        for line in text.split(
            "\n"
        ):

            line = line.strip()

            if not line:
                continue

            # ----------------------------------------------
            # Long lines
            # ----------------------------------------------

            if len(line) > 500:

                sentence_parts = re.split(
                    r"(?<=[.!?])\s+",
                    line,
                )

                parts.extend(
                    sentence_parts
                )

            else:

                parts.append(
                    line
                )

        sentences = parts


    # --------------------------------------------------------
    # Final sentence cleanup
    # --------------------------------------------------------

    sentences = [
        sentence.strip()
        for sentence in sentences
        if sentence
        and sentence.strip()
    ]


    if not sentences:
        return []


    # ========================================================
    # BUILD CHUNKS
    # ========================================================

    chunks: List[str] = []

    current_tokens: List[str] = []

    for sentence in sentences:

        tokens = sentence.split()

        if not tokens:
            continue

        current_tokens.extend(
            tokens
        )

        while (
            len(current_tokens)
            >= chunk_size
        ):

            chunk = " ".join(
                current_tokens[
                    :chunk_size
                ]
            ).strip()

            if chunk:
                chunks.append(
                    chunk
                )

            # ------------------------------------------------
            # Keep overlap from previous chunk.
            # ------------------------------------------------

            current_tokens = (
                current_tokens[
                    chunk_size - overlap:
                ]
            )


    # ========================================================
    # FINAL PARTIAL CHUNK
    # ========================================================

    if current_tokens:

        final_chunk = " ".join(
            current_tokens
        ).strip()

        if final_chunk:
            chunks.append(
                final_chunk
            )


    # ========================================================
    # REMOVE DUPLICATE CHUNKS
    # ========================================================

    seen = set()

    unique_chunks = []

    for chunk in chunks:

        normalized = chunk.strip()

        if (
            normalized
            and normalized not in seen
        ):

            unique_chunks.append(
                normalized
            )

            seen.add(
                normalized
            )

    return unique_chunks